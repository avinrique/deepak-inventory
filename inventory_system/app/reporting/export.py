"""Turns a ReportResult into CSV/Excel/PDF files, or sends it to a printer.

Written against the standard library's csv module and openpyxl directly.
This used to build a pandas DataFrame first, which brought pandas and numpy
into the installer — roughly 90 MB and a second of import time — to do four
things pandas was never needed for: no aggregation happens here, every row
arrives pre-computed from app.repositories.sql.reporting_repository, and the
quoting/encoding/.xlsx work is what csv and openpyxl already do.

PDF export and the Print action both render the same HTML table through
Qt's own QTextDocument + QtPrintSupport (QPrinter/QPrintDialog) — already a
project dependency (PySide6), so no reportlab/wkhtmltopdf install needed,
and "export to PDF" and "print" share one rendering path instead of two.

"""
import csv as csv_lib
import html as html_lib
from datetime import date, datetime
from decimal import Decimal

from app.schemas.reporting import ReportResult


_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _neutralize_formula(value: str) -> str:
    """Defuses CSV/Excel formula injection (CWE-1236): a cell value that
    starts with =, +, -, @, tab, or CR is interpreted as a formula by
    Excel/LibreOffice on open. Report cells often echo free-text user input
    (customer/supplier/product names, notes), so any of those must be
    neutralized before writing, not just PDF/HTML cells (already safe via
    html.escape).
    """
    if value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


def _format_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:,}"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _coerce(value):
    """Decimal -> float at this boundary only: spreadsheet cells need to be
    numeric (so a user can sum/format them in Excel), and the underlying
    calculation was already done in Decimal upstream — this conversion is
    purely for display, not further financial computation.
    """
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str):
        return _neutralize_formula(value)
    return value


def to_rows(result: ReportResult) -> list[list]:
    """Rows as flat lists in `result.columns` order, coerced for a
    spreadsheet. A column a row does not carry becomes an empty cell rather
    than a KeyError — reports assemble their rows per query, and a missing
    optional column should not fail the export."""
    return [[_coerce(row.get(column)) for column in result.columns]
            for row in result.rows]


def export_csv(result: ReportResult, path: str) -> str:
    # newline="" is required by the csv module: without it, its own \r\n
    # line endings get translated again on Windows, giving every row a blank
    # line after it.
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv_lib.writer(handle)
        writer.writerow(result.columns)
        writer.writerows(to_rows(result))
    return path


def export_excel(result: ReportResult, path: str) -> str:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (result.title or "Report")[:31]  # Excel's sheet-name length limit
    sheet.append(list(result.columns))
    for row in to_rows(result):
        sheet.append(row)
    workbook.save(path)
    return path


def render_html_table(result: ReportResult) -> str:
    header_cells = "".join(f"<th>{html_lib.escape(str(col))}</th>" for col in result.columns)
    body_rows = []
    for row in result.rows:
        cells = "".join(f"<td>{html_lib.escape(_format_cell(row.get(col)))}</td>"
                        for col in result.columns)
        body_rows.append(f"<tr>{cells}</tr>")
    generated = result.generated_at.strftime("%Y-%m-%d %H:%M")
    return f"""<html><head><style>
        body {{ font-family: sans-serif; font-size: 11px; }}
        h2 {{ margin-bottom: 2px; }}
        .meta {{ color: #666; font-size: 10px; margin-bottom: 12px; }}
        table {{ border-collapse: collapse; width: 100%; }}
        th, td {{ border: 1px solid #ccc; padding: 4px 8px; text-align: left; }}
        th {{ background: #f0f0f0; }}
    </style></head>
    <body>
        <h2>{html_lib.escape(result.title)}</h2>
        <div class="meta">Generated {generated} &mdash; {result.row_count} row(s)</div>
        <table><thead><tr>{header_cells}</tr></thead><tbody>{"".join(body_rows)}</tbody></table>
    </body></html>"""


def export_pdf(result: ReportResult, path: str) -> str:
    from PySide6.QtGui import QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    document = QTextDocument()
    document.setHtml(render_html_table(result))

    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(path)
    document.print_(printer)
    return path


def print_report(result: ReportResult, parent=None) -> bool:
    """Shows the OS print dialog; returns False if the user cancels it
    rather than raising, since cancelling is a normal outcome here, not an
    error condition.
    """
    from PySide6.QtGui import QTextDocument
    from PySide6.QtPrintSupport import QPrintDialog, QPrinter

    document = QTextDocument()
    document.setHtml(render_html_table(result))

    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    dialog = QPrintDialog(printer, parent)
    if dialog.exec() != QPrintDialog.DialogCode.Accepted:
        return False
    document.print_(printer)
    return True
